#!/usr/bin/env python3
"""NanoHIV-DR clinical report generator v2.

Normalises current Stanford HIVdb consensus + By-Reads JSON into one sample
object before rendering a concise four-section clinical DOCX report.
"""
import argparse, csv, datetime, json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

SCRIPT_DIR = Path(__file__).resolve().parent
REPORT_DIR = SCRIPT_DIR.parent
LOGO1 = REPORT_DIR / "logoUVRI.png"
LOGO2 = REPORT_DIR / "logoCVR.png"

import re as _re

CODON_TABLE = {
    'TTT':'F','TTC':'F','TTA':'L','TTG':'L','CTT':'L','CTC':'L','CTA':'L','CTG':'L',
    'ATT':'I','ATC':'I','ATA':'I','ATG':'M','GTT':'V','GTC':'V','GTA':'V','GTG':'V',
    'TCT':'S','TCC':'S','TCA':'S','TCG':'S','CCT':'P','CCC':'P','CCA':'P','CCG':'P',
    'ACT':'T','ACC':'T','ACA':'T','ACG':'T','GCT':'A','GCC':'A','GCA':'A','GCG':'A',
    'TAT':'Y','TAC':'Y','TAA':'*','TAG':'*','CAT':'H','CAC':'H','CAA':'Q','CAG':'Q',
    'AAT':'N','AAC':'N','AAA':'K','AAG':'K','GAT':'D','GAC':'D','GAA':'E','GAG':'E',
    'TGT':'C','TGC':'C','TGA':'*','TGG':'W','CGT':'R','CGC':'R','CGA':'R','CGG':'R',
    'AGT':'S','AGC':'S','AGA':'R','AGG':'R','GGT':'G','GGC':'G','GGA':'G','GGG':'G',
}

_MUT_RE = _re.compile(r'^[A-Za-z\*]*?(\d+)([A-Za-z\*]+)$')
_AA = set("ACDEFGHIKLMNPQRSTVWY*")


def _translate(codon):
    codon = (codon or "").upper()
    if len(codon) != 3 or "-" in codon or "N" in codon:
        return None
    return CODON_TABLE.get(codon)


def load_codfreq(path):
    """Return {'counts': {(gene,pos): {aa: n}}, 'totals': {(gene,pos): n}} or None."""
    counts, totals = {}, {}
    try:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                try:
                    gene = row["gene"]
                    pos = int(row["position"])
                    total = int(row["total"])
                    count = int(row["count"])
                except (KeyError, ValueError, TypeError):
                    continue
                totals[(gene, pos)] = total
                aa = _translate(row.get("codon", ""))
                if aa is None:
                    continue
                counts.setdefault((gene, pos), {})[aa] = counts.get((gene, pos), {}).get(aa, 0) + count
    except OSError:
        return None
    if not totals:
        return None
    return {"counts": counts, "totals": totals}


def format_reads_prevalence(prevmap, gene, name):
    """Prevalence of a mutation's alt amino acid(s) from CodFreq counts, e.g. '99.5%'."""
    if not prevmap:
        return "n/a"
    raw = name.strip()
    m = _MUT_RE.match(raw)
    if not m:
        return "n/a"
    alt_raw = m.group(2)
    if alt_raw != alt_raw.upper():   # lowercase => insertion/deletion notation
        return "n/a"
    pos = int(m.group(1))
    alt = [a for a in alt_raw.upper() if a in _AA]
    if not alt:
        return "n/a"
    total = prevmap["totals"].get((gene, pos))
    if not total:
        return "n/a"
    d = prevmap["counts"].get((gene, pos), {})
    c = sum(d.get(a, 0) for a in alt)
    return f"{100.0 * c / total:.1f}%"


GENE_NAMES = {"PR": "Protease", "RT": "Reverse Transcriptase", "IN": "Integrase"}
CLASS_NAMES = {
    "PI": "Protease Inhibitors",
    "NRTI": "Nucleoside Reverse Transcriptase Inhibitors",
    "NNRTI": "Non-Nucleoside Reverse Transcriptase Inhibitors",
    "INSTI": "Integrase Strand Transfer Inhibitors",
}
CLASS_ORDER = ["PI", "NRTI", "NNRTI", "INSTI"]


def as_record(value):
    if isinstance(value, list):
        return value[0] if value else {}
    return value if isinstance(value, dict) else {}


def safe_text(value, default="Not provided"):
    if value is None or str(value).strip() == "":
        return default
    return str(value).strip()


def fmt_int(value):
    return f"{value:,.0f}" if isinstance(value, (int, float)) else "NA"


def fmt_fraction_pct(value):
    """Format a 0-1 fraction (e.g. mixtureRate) as a percentage."""
    return f"{value * 100:.1f}%" if isinstance(value, (int, float)) else "NA"


def fmt_hivdb_prevalence(value):
    """Format Sierra By-Reads prevalence thresholds stored as 0-1 fractions."""
    if not isinstance(value, (int, float)):
        return "NA"
    return f"{value * 100:.1f}%"


def set_cell_text(cell, text, bold=False, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def load_metadata(path):
    out = {}
    with open(path, newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        if "sample_id" not in (reader.fieldnames or []):
            raise ValueError("Metadata must contain a sample_id column")
        for row in reader:
            sid = (row.get("sample_id") or "").strip()
            if sid:
                out[sid] = {k: (v or "").strip() for k, v in row.items()}
    return out


def load_reads_dir(path):
    if not path:
        return {}
    directory = Path(path)
    if not directory.is_dir():
        raise ValueError(f"By-Reads directory does not exist: {directory}")
    out = {}
    for f in sorted(directory.glob("*.reads.json")):
        with open(f, encoding="utf-8") as fh:
            out[f.name.replace(".reads.json", "")] = as_record(json.load(fh))
    return out


def extract_drugs(record):
    drugs = {}
    for gene_result in record.get("drugResistance", []) or []:
        gene = gene_result.get("gene", {}).get("name", "Unknown")
        for ds in gene_result.get("drugScores", []) or []:
            drug = ds.get("drug", {}) or {}
            name = drug.get("displayAbbr") or drug.get("name") or "Unknown"
            klass = ds.get("drugClass", {}).get("name", "Unknown")
            drugs[name] = {
                "drug": name,
                "gene": gene,
                "class": klass,
                "level": ds.get("level"),
                "score": ds.get("score"),
                "text": ds.get("text") or "Not interpreted",
            }
    return drugs


def extract_mutations_and_comments(record):
    mutations = {}
    # All called amino-acid mutations from aligned gene sequences.
    for gs in record.get("alignedGeneSequences", []) or []:
        gene = gs.get("gene", {}).get("name", "Unknown")
        for m in gs.get("mutations", []) or []:
            if not isinstance(m, dict):
                continue
            text = m.get("text") or m.get("displayAAs") or m.get("AAs")
            if text:
                mutations.setdefault(str(text), {"gene": gene, "classes": set(), "comments": set()})

    # Resistance-associated mutations and Stanford comments from partial scores.
    for dr in record.get("drugResistance", []) or []:
        gene = dr.get("gene", {}).get("name", "Unknown")
        for ds in dr.get("drugScores", []) or []:
            klass = ds.get("drugClass", {}).get("name", "")
            for partial in ds.get("partialScores", []) or []:
                for m in partial.get("mutations", []) or []:
                    text = m.get("text", "")
                    if not text:
                        continue
                    entry = mutations.setdefault(text, {"gene": gene, "classes": set(), "comments": set()})
                    if klass:
                        entry["classes"].add(klass)
                    if m.get("primaryType"):
                        entry["classes"].add(m["primaryType"])
                    for c in m.get("comments", []) or []:
                        if c.get("text"):
                            entry["comments"].add(c["text"].strip())
    return mutations


def extract_coverage(record):
    coverage = []
    for gs in record.get("alignedGeneSequences", []) or []:
        gene = gs.get("gene", {}).get("name", "Unknown")
        coverage.append({"gene": gene, "first": gs.get("firstAA"), "last": gs.get("lastAA")})
    return coverage


def mutation_label(m):
    """Build a Stanford-style amino-acid mutation label when one is not supplied."""
    if not isinstance(m, dict):
        return None
    text = m.get("text") or m.get("mutation") or m.get("display")
    if text:
        return str(text)
    ref, pos, aas = m.get("reference"), m.get("position"), m.get("AAs")
    if ref and pos is not None and aas:
        return f"{ref}{pos}{aas}"
    return None


def extract_reads_mutations(reads):
    """Return mutations detected by Sierra By-Reads without inventing prevalence.

    Current Sierra allGeneSequenceReads mutation objects identify reference,
    position and called amino acid(s), but do not carry a per-mutation prevalence.
    Resistance partialScores independently identify the mutations used by HIVdb.
    """
    found = {}
    for gs in reads.get("allGeneSequenceReads", []) or []:
        if not isinstance(gs, dict):
            continue
        gene = (gs.get("gene") or {}).get("name", "Unknown")
        for m in gs.get("mutations", []) or []:
            label = mutation_label(m)
            if label:
                found[label] = {"gene": gene}
    for dr in reads.get("drugResistance", []) or []:
        gene = (dr.get("gene") or {}).get("name", "Unknown")
        for ds in dr.get("drugScores", []) or []:
            for partial in ds.get("partialScores", []) or []:
                for m in partial.get("mutations", []) or []:
                    label = mutation_label(m)
                    if label:
                        found.setdefault(label, {"gene": gene})
    return found


def extract_reads_coverage(reads):
    coverage = []
    for gs in reads.get("allGeneSequenceReads", []) or []:
        if not isinstance(gs, dict):
            continue
        gene = (gs.get("gene") or {}).get("name", "Unknown")
        coverage.append({"gene": gene, "first": gs.get("firstAA"), "last": gs.get("lastAA")})
    return coverage


def assay_qc_status(qc):
    """Assess depth QC at drug-resistance positions within the sequenced assay region."""
    threshold = qc.get("min_position_reads")
    drp_min = (qc.get("depth_drp") or {}).get("min")
    try:
        if threshold is None or drp_min is None:
            return "REVIEW"
        return "PASS" if float(drp_min) >= float(threshold) else "REVIEW"
    except (TypeError, ValueError):
        return "REVIEW"


def concordance(consensus, reads):
    if consensus is None or reads is None:
        return "Not assessed"
    cl, rl = consensus.get("level"), reads.get("level")
    if isinstance(cl, (int, float)) and isinstance(rl, (int, float)):
        if cl == rl:
            return "Concordant"
        return "By-Reads higher" if rl > cl else "By-Reads lower"
    if consensus.get("text") == reads.get("text"):
        return "Concordant"
    return "Discordant"



def interpretation_concordance_status(consensus_drugs, reads_drugs):
    """Summarise consensus vs By-Reads drug-resistance agreement."""
    all_drugs = set(consensus_drugs) | set(reads_drugs)
    if not all_drugs:
        return "NOT ASSESSED"
    for name in all_drugs:
        if concordance(consensus_drugs.get(name), reads_drugs.get(name)) != "Concordant":
            return "REVIEW REQUIRED"
    return "CONCORDANT"


def discordant_drug_count(consensus_drugs, reads_drugs):
    all_drugs = set(consensus_drugs) | set(reads_drugs)
    return sum(
        concordance(consensus_drugs.get(name), reads_drugs.get(name)) != "Concordant"
        for name in all_drugs
    )

def normalise_sample(consensus, reads, metadata):
    sample = consensus.get("inputSequence", {}).get("header", "").strip()
    reads = as_record(reads)
    depth = reads.get("readDepthStats", {}) or {}
    validation = []
    for x in reads.get("validationResults", []) or []:
        msg = x.get("message", "") if isinstance(x, dict) else str(x)
        level = x.get("level", "") if isinstance(x, dict) else ""
        if msg:
            validation.append({"level": level, "message": msg})

    cdrugs = extract_drugs(consensus)
    rdrugs = extract_drugs(reads)
    mutations = extract_mutations_and_comments(consensus)

    # Mark mutations originating from the consensus interpretation explicitly.
    for entry in mutations.values():
        entry["consensus_detected"] = True
        entry.setdefault("reads_detected", False)

    # Build the union of resistance-associated mutations from consensus and By-Reads.
    # By-Reads partialScores provide the resistance class/comments needed to recognise
    # clinically relevant mutations even when they are absent from consensus.
    reads_mutations = extract_mutations_and_comments(reads)
    for name, info in reads_mutations.items():
        if not (info.get("classes") or info.get("comments")):
            continue
        entry = mutations.setdefault(
            name,
            {
                "gene": info.get("gene", "Unknown"),
                "classes": set(),
                "comments": set(),
                "consensus_detected": False,
                "reads_detected": False,
            },
        )
        entry["reads_detected"] = True
        if entry.get("gene") in ("", "Unknown"):
            entry["gene"] = info.get("gene", "Unknown")
        entry["classes"].update(info.get("classes", set()))
        entry["comments"].update(info.get("comments", set()))

    subtype_reads = reads.get("bestMatchingSubtype", {}) or {}
    if isinstance(subtype_reads, dict):
        subtype_reads = subtype_reads.get("display") or subtype_reads.get("name") or "Not determined"

    return {
        "sample_id": sample,
        "metadata": metadata,
        "subtype_consensus": consensus.get("subtypeText") or "Not determined",
        "subtype_reads": subtype_reads or "Not determined",
        "coverage": extract_coverage(consensus),
        "reads_coverage": extract_reads_coverage(reads),
        "consensus_drugs": cdrugs,
        "reads_drugs": rdrugs,
        "mutations": mutations,
        "qc": {
            "configured_min_prevalence": reads.get("minPrevalence"),
            "actual_min_prevalence": reads.get("actualMinPrevalence"),
            "min_codon_reads": reads.get("minCodonReads"),
            "min_position_reads": reads.get("minPositionReads"),
            "mixture_rate": reads.get("mixtureRate"),
            "depth": depth,
            "depth_drp": reads.get("readDepthStatsDRP", {}) or {},
            "validation": validation,
        },
    }


def configure_document(doc, sample_id):
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.7)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)

    # Repeating header on every page: laboratory logo on the left,
    # explicit sample identifier on the right.
    header = sec.header
    p = header.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_after = Pt(0)

    header_table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    header_table.autofit = False
    left_cell, right_cell = header_table.rows[0].cells
    left_cell.width = Inches(4.6)
    right_cell.width = Inches(1.9)

    left_p = left_cell.paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if LOGO1.exists():
        left_p.add_run().add_picture(str(LOGO1), width=Inches(2.2))

    right_p = right_cell.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sample_run = right_p.add_run(f"Sample ID: {sample_id}")
    sample_run.bold = True
    sample_run.font.size = Pt(10)

    add_page_number(sec.footer.paragraphs[0])


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    if subtitle:
        r = doc.add_paragraph(subtitle)
        r.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], safe_text(value))
    return table


def page1(doc, s):
    add_title(doc, "HIV-1 Genotypic Drug Resistance Report")
    doc.add_heading("Sample information", level=1)
    m = s["metadata"]
    add_key_value_table(doc, [
        ("Sample ID", s["sample_id"]),
        ("Patient ID", m.get("patient_id")),
        ("Age", m.get("age")),
        ("Sex", m.get("sex")),
        ("Sample collection date", m.get("date")),
    ])

    doc.add_heading("Assay summary", level=1)
    rows = [
        ("Consensus subtype", s["subtype_consensus"]),
        ("By-Reads subtype", s["subtype_reads"]),
    ]
    for c in s["coverage"]:
        rows.append((f'Consensus {c["gene"]} coverage', f'Codons {c["first"]}-{c["last"]}'))
    for c in s["reads_coverage"]:
        rows.append((f'By-Reads {c["gene"]} coverage', f'Codons {c["first"]}-{c["last"]}'))
    d = s["qc"]["depth"]
    rows.extend([
        ("Median read depth", f'{fmt_int(d.get("p50"))}x'),
        ("Configured minimum variant prevalence", fmt_hivdb_prevalence(s["qc"]["configured_min_prevalence"])),
        ("Target-region QC", assay_qc_status(s["qc"])),
        ("Consensus/By-Reads concordance",
         interpretation_concordance_status(s["consensus_drugs"], s["reads_drugs"])),
    ])
    add_key_value_table(doc, rows)

    doc.add_paragraph(
        "The percentage shown with each subtype is the Stanford HIVdb nucleotide distance to the "
        "closest reference sequence (a lower value indicates a closer match); it is not a confidence score."
    )

    n_discordant = discordant_drug_count(s["consensus_drugs"], s["reads_drugs"])
    if n_discordant:
        p = doc.add_paragraph()
        r = p.add_run(
            f"REVIEW REQUIRED — {n_discordant} drug-resistance interpretation"
            f"{'s are' if n_discordant != 1 else ' is'} discordant between consensus and By-Reads."
        )
        r.bold = True

    doc.add_paragraph("Drug-resistance interpretations are generated using Stanford HIVdb. Results should be interpreted with sequence coverage and assay QC information.")


def page2(doc, s):
    doc.add_page_break()
    doc.add_heading("HIV Drug Resistance Results", level=1)
    doc.add_paragraph("Consensus and Stanford HIVdb By-Reads interpretations are shown side by side. Concordance is based on the Stanford resistance level when available.")
    all_drugs = set(s["consensus_drugs"]) | set(s["reads_drugs"])
    for klass in CLASS_ORDER + ["Unknown"]:
        names = [n for n in all_drugs if (s["consensus_drugs"].get(n) or s["reads_drugs"].get(n) or {}).get("class", "Unknown") == klass]
        if not names:
            continue
        doc.add_heading(CLASS_NAMES.get(klass, klass), level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for c, text in zip(table.rows[0].cells, ["Drug", "Consensus", "By-Reads", "Concordance"]):
            set_cell_text(c, text, bold=True)
            shade_cell(c, "D9EAF7")
        for name in sorted(names):
            c = s["consensus_drugs"].get(name)
            r = s["reads_drugs"].get(name)
            cells = table.add_row().cells
            vals = [name, c.get("text") if c else "Not assessed", r.get("text") if r else "Not assessed", concordance(c, r)]
            for cell, value in zip(cells, vals):
                set_cell_text(cell, value)


def page3(doc, s):
    doc.add_page_break()
    doc.add_heading("Resistance-Associated Mutations", level=1)
    resistance = [(name, x) for name, x in s["mutations"].items() if x.get("classes") or x.get("comments")]
    if not resistance:
        doc.add_paragraph("No resistance-associated mutations were identified in either the consensus or By-Reads Stanford HIVdb interpretation.")
        return
    prevmap = s.get("reads_prevalence")
    have_prev = bool(prevmap)
    headers = ["Gene", "Mutation", "Class", "Consensus", "By-Reads"]
    if have_prev:
        headers.append("By-Reads prevalence")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for c, text in zip(table.rows[0].cells, headers):
        set_cell_text(c, text, bold=True)
        shade_cell(c, "D9EAF7")
    for name, x in sorted(resistance, key=lambda z: (z[1].get("gene", ""), z[0])):
        vals = [
            x.get("gene", ""),
            name,
            ", ".join(sorted(x.get("classes", []))) or "Resistance-associated",
            "Detected" if x.get("consensus_detected") else "Not detected",
            "Detected" if x.get("reads_detected") else "Not detected",
        ]
        if have_prev:
            vals.append(format_reads_prevalence(prevmap, x.get("gene", ""), name))
        cells = table.add_row().cells
        for cell, value in zip(cells, vals):
            set_cell_text(cell, value)

    if have_prev:
        doc.add_paragraph(
            "By-Reads prevalence is the proportion of reads supporting the mutant amino acid at that "
            "codon, computed directly from the CodFreq read counts. \"n/a\" indicates a position that "
            "could not be resolved unambiguously (for example an insertion, deletion, or unmatched codon)."
        )
    else:
        doc.add_paragraph(
            "By-Reads indicates whether the mutation was identified in the supplied Stanford HIVdb "
            "By-Reads result. Individual mutation prevalence is not reported because CodFreq read "
            "counts were not supplied to the report generator."
        )

    comments = [(name, x) for name, x in resistance if x.get("comments")]
    if comments:
        doc.add_heading("Stanford HIVdb interpretation/comments", level=1)
        seen = set()
        for name, x in comments:
            unique_comments = [c for c in sorted(x["comments"]) if c not in seen]
            for comment in unique_comments:
                p = doc.add_paragraph()
                p.add_run(name + ": ").bold = True
                p.add_run(comment)
                seen.add(comment)


def page4(doc, s):
    doc.add_page_break()
    doc.add_heading("Technical and Quality-Control Information", level=1)
    d = s["qc"]["depth"]
    drp = s["qc"]["depth_drp"]
    add_key_value_table(doc, [
        ("Mean read depth", f'{fmt_int(d.get("mean"))}x'),
        ("Median read depth", f'{fmt_int(d.get("p50"))}x'),
        ("Minimum read depth", f'{fmt_int(d.get("min"))}x'),
        ("Maximum read depth", f'{fmt_int(d.get("max"))}x'),
        ("Median depth at drug-resistance positions", f'{fmt_int(drp.get("p50"))}x'),
        ("Minimum depth at drug-resistance positions", f'{fmt_int(drp.get("min"))}x'),
        ("Configured minimum variant prevalence", fmt_hivdb_prevalence(s["qc"]["configured_min_prevalence"])),
        ("Effective HIVdb variant-calling threshold", fmt_hivdb_prevalence(s["qc"]["actual_min_prevalence"])),
        ("Minimum position read depth", fmt_int(s["qc"]["min_position_reads"])),
        ("Minimum codon reads", fmt_int(s["qc"]["min_codon_reads"])),
        ("Mixture rate", fmt_fraction_pct(s["qc"]["mixture_rate"])),
    ])
    doc.add_paragraph(
        "Stanford HIVdb By-Reads prevalence thresholds are stored as proportions and are shown here as percentages. "
        "The configured minimum prevalence is the requested calling threshold. The effective HIVdb "
        "variant-calling threshold is the sample-specific threshold applied by the By-Reads analysis; "
        "it is not the prevalence of any individual resistance mutation."
    )

    doc.add_heading("Sequence coverage", level=2)
    if s["reads_coverage"]:
        for c in s["reads_coverage"]:
            doc.add_paragraph(
                f'By-Reads {GENE_NAMES.get(c["gene"], c["gene"])} ({c["gene"]}): '
                f'codons {c["first"]}-{c["last"]}', style="List Bullet"
            )
    else:
        for c in s["coverage"]:
            doc.add_paragraph(
                f'{GENE_NAMES.get(c["gene"], c["gene"])} ({c["gene"]}): codons {c["first"]}-{c["last"]}',
                style="List Bullet",
            )

    qc_status = assay_qc_status(s["qc"])
    _qc_genes = [c["gene"] for c in (s.get("reads_coverage") or s.get("coverage") or [])]
    region_label = "/".join(_qc_genes) if _qc_genes else "target"
    if qc_status == "PASS":
        doc.add_paragraph(
            "Target-region QC: PASS — all evaluated drug-resistance positions within the targeted "
            f"{region_label} region met the configured minimum read-depth requirement."
        )
    else:
        doc.add_paragraph(
            "Target-region QC: REVIEW — depth at one or more evaluated drug-resistance positions "
            f"within the targeted {region_label} region did not meet, or could not be compared with, the "
            "configured minimum read-depth requirement."
        )

    doc.add_heading("Additional Stanford HIVdb validation messages", level=2)
    if s["qc"]["validation"]:
        for v in s["qc"]["validation"]:
            prefix = f'{v["level"]}: ' if v["level"] else ""
            doc.add_paragraph(prefix + v["message"], style="List Bullet")
        doc.add_paragraph(
            "Coverage note: Stanford HIVdb may report unsequenced resistance positions outside "
            f"the targeted {region_label} amplicon. These positions do not represent coverage failures within "
            "the intended assay region and do not, by themselves, cause the target-region QC assessment to fail."
        )
    else:
        doc.add_paragraph("No Stanford HIVdb validation warnings were reported.")


def build_report(sample, output):
    doc = Document()
    configure_document(doc, sample["sample_id"])
    page1(doc, sample)
    page2(doc, sample)
    page3(doc, sample)
    page4(doc, sample)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description="NanoHIV-DR clinical report generator v2")
    parser.add_argument("--json", required=True, help="Consensus Stanford HIVdb JSON")
    parser.add_argument("--data", required=True, help="Tab-delimited metadata")
    parser.add_argument("--output", default="subtypes.tsv", help="Subtype summary TSV")
    parser.add_argument("--reports", action="store_true", help="Generate DOCX reports")
    parser.add_argument("--reads-dir", help="Directory containing *.reads.json")
    parser.add_argument("--fasta", help="Optional consensus FASTA (reserved for future callability QC)")
    parser.add_argument("--report-dir", help="DOCX output directory")
    parser.add_argument("--codfreq-dir", help="Directory of *.codfreq files for per-mutation prevalence")
    args = parser.parse_args()

    metadata = load_metadata(args.data)
    reads = load_reads_dir(args.reads_dir)
    codfreq_files = {}
    if args.codfreq_dir:
        for cf in Path(args.codfreq_dir).glob("*.codfreq"):
            sid = cf.name
            for suf in ("_filtered.codfreq", ".codfreq"):
                if sid.endswith(suf):
                    sid = sid[: -len(suf)]
                    break
            codfreq_files[sid] = cf
    with open(args.json, encoding="utf-8") as fh:
        consensus = json.load(fh)
    if not isinstance(consensus, list):
        consensus = [consensus]

    report_dir = Path(args.report_dir or (datetime.datetime.now().strftime("%d-%m-%Y") + "_reports"))
    if args.reports:
        report_dir.mkdir(parents=True, exist_ok=True)

    subtypes = []
    for rec in consensus:
        sid = rec.get("inputSequence", {}).get("header", "").strip()
        if not sid:
            print("WARNING: consensus record has no inputSequence.header; skipping")
            continue
        if sid not in metadata:
            print(f"WARNING: no metadata for sample '{sid}'; skipping")
            continue
        s = normalise_sample(rec, reads.get(sid), metadata[sid])
        if sid in codfreq_files:
            s["reads_prevalence"] = load_codfreq(codfreq_files[sid])
        subtypes.append((sid, s["subtype_consensus"]))
        print(f"Processing sample '{sid}' (Patient ID: {safe_text(metadata[sid].get('patient_id'), 'not provided')})")
        if args.reports:
            build_report(s, report_dir / f"{sid}_report.docx")

    with open(args.output, "w", encoding="utf-8") as out:
        for sid, subtype in subtypes:
            out.write(f"{sid}\t{subtype}\n")
    print(f"Generated {len(subtypes)} sample result(s)")
    if args.reports:
        print(f"Clinical reports: {report_dir}")


if __name__ == "__main__":
    main()
